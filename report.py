"""
REPORT EXPORT: the dashboard as a PDF, a Word file, an Excel file and a web page.

No Streamlit in this file, so every number can be tested. The numbers come
from builder.compute_tile (the same code that draws the screen), so the
report always shows exactly what the dashboard shows.

    data = build_report(engine, info, tiles, filters, opts, title)
    to_pdf(data), to_docx(data), to_xlsx(data), to_html(data)   -> bytes / text
"""

from safe_numbers import is_whole
import datetime
import html as html_lib
import importlib
import io
import re

import builder as bd
import insights as ins
import ui

LIBRARIES = {
    "vl_convert": "vl-convert-python",
    "docx": "python-docx",
    "reportlab": "reportlab",
    "openpyxl": "openpyxl",
}

FORMATS = {
    "pdf": ("PDF", "pdf", "application/pdf"),
    "docx": ("Word", "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "xlsx": ("Excel", "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "html": ("Web page", "html", "text/html"),
}

CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def missing_libraries():
    """The pip names of the libraries that are not installed."""

    missing = []

    for module, pip_name in LIBRARIES.items():
        try:
            importlib.import_module(module)
        except Exception:
            missing.append(pip_name)

    return missing


def clean(text):
    """Text that every file format accepts (no control characters)."""

    return CONTROL_CHARS.sub("", str(text))


def file_name(title, extension):
    slug = re.sub(r"[^A-Za-z0-9_-]+", "_", clean(title)).strip("_")[:50] or "report"
    return f"{slug}.{extension}"


# ============================================================
# The report data (the numbers)
# ============================================================

def describe_filters(info, filters, opts):
    """Plain sentences about what the report is limited to."""

    lines = []

    if filters.get("months"):
        lines.append(f"Months: {filters['months'][0]} to {filters['months'][1]}")

    for column, chosen in (filters.get("values") or {}).items():

        label = info["labels"].get(column, column).capitalize()

        if not chosen:
            lines.append(f"{label}: nothing selected")
            continue

        shown = ", ".join(str(v) for v in chosen[:6])

        if len(chosen) > 6:
            shown += f" and {len(chosen) - 6} more"

        lines.append(f"{label}: {shown}")

    if info["dates"]:
        lines.append(f"Time view: {bd.GRAINS[opts.get('grain') or 'M']}")

    if opts.get("top"):
        lines.append(f"Show top: {opts['top']}")

    return lines


def build_report(engine, info, tiles, filters, opts, title="", file_name_="", with_insights=True):
    """Everything the four exporters need. Never raises for odd dashboards."""

    opts = opts or {"grain": "M", "top": None}
    view = bd.apply_filters(info, filters or {})

    kpis, charts = [], []

    for tile in tiles:

        result = bd.compute_tile(view, tile, info, opts)
        name = tile.get("title") or bd.auto_title(tile, info, opts.get("grain") or "M")

        if tile["type"] == "kpi":

            if result["kind"] == "kpi" and result["value"] is not None:
                kpis.append({
                    "label": name, "value": bd.kpi_number(result["value"]),
                    "raw": float(result["value"]),
                    "sub": f"{result['rows']:,} {info['noun']} in view",
                })
            else:
                kpis.append({
                    "label": name, "value": "—", "raw": None,
                    "sub": result.get("message", "No numbers to calculate"),
                })

        else:
            charts.append({
                "title": name, "type": tile["type"], "result": result,
                "subtitle": bd.tile_subtitle(tile, result),
            })

    df = info["df"]

    return {
        "title": clean(title).strip() or "Dashboard report",
        "file": clean(file_name_),
        "created": datetime.datetime.now().strftime("%d %B %Y, %H:%M"),
        "noun": info["noun"],
        "rows_in_view": len(view),
        "rows_total": len(df),
        "filters": describe_filters(info, filters or {}, opts),
        "kpis": kpis,
        "charts": charts,
        "insights": ins.build_insights(engine) if with_insights else [],
        "quality": {
            "rows": len(df), "columns": len(df.columns),
            "missing": int(df.isnull().sum().sum()),
            "duplicates": int(df.duplicated().sum()),
        },
    }


# ============================================================
# Charts as pictures (drawn without a browser)
# ============================================================

def chart_of(item, width=560):
    """The Altair chart for a report item, or None (a table / no data)."""

    result = item["result"]

    if result["kind"] == "line":

        marked = {
            (f"{k}*" if k in result["partial"] else k): v
            for k, v in result["months"].items()
        }
        chart = ui.line_chart(marked, height=300, export=True)

    elif result["kind"] == "groups" and item["type"] != "table":

        values = result["values"]

        if item["type"] == "donut":
            chart = ui.donut_chart(values, height=300)
        elif item["type"] == "column":
            chart = ui.column_chart(values, height=300)
        else:
            chart = ui.bar_chart(values, height=max(220, 34 * len(values) + 50))

    else:
        return None

    # a solid white background (transparent pictures lose their colours in a PDF)
    return chart.properties(width=width, background="white")


def png_of(item):
    """PNG bytes of the chart (made once, then remembered), or None."""

    if "png" not in item:

        import vl_convert

        chart = chart_of(item)
        item["png"] = (
            vl_convert.vegalite_to_png(chart.to_json(), scale=2) if chart else None
        )

    return item["png"]


def svg_of(item, number):
    """SVG text of the chart with ids that cannot clash with other charts."""

    import vl_convert

    chart = chart_of(item, width=760)

    if chart is None:
        return None

    svg = vl_convert.vegalite_to_svg(chart.to_json())
    prefix = f"c{number}_"

    svg = re.sub(r'\bid="([^"]+)"', lambda m: f'id="{prefix}{m.group(1)}"', svg)
    svg = re.sub(r"url\(#([^)]+)\)", lambda m: f"url(#{prefix}{m.group(1)})", svg)
    svg = re.sub(r'(xlink:)?href="#([^"]+)"', lambda m: f'{m.group(1) or ""}href="#{prefix}{m.group(2)}"', svg)

    return svg


def rows_of(item):
    """(header, rows) of the numbers behind a chart: [(name, value, share or None)]."""

    result = item["result"]

    if result["kind"] == "line":
        return ["Period", "Value"], [(k, v, None) for k, v in result["months"].items()]

    if result["kind"] == "groups":

        frame = result.get("frame")

        if frame is not None and "Share %" in frame:
            shares = list(frame["Share %"])
        else:
            shares = [None] * len(result["values"])

        header = ["Group", "Value"] + (["Share %"] if frame is not None and "Share %" in frame else [])

        return header, [
            (k, v, s) for (k, v), s in zip(result["values"].items(), shares)
        ]

    return None, []


def number_text(value):
    if value is None:
        return ""
    if is_whole(float(value)):
        return f"{int(value):,}"
    return f"{value:,.2f}"


# ============================================================
# PDF
# ============================================================

def pdf_plain(text):
    """Text the standard PDF fonts can draw (Rs for the rupee sign)."""

    text = clean(text).replace("₹", "Rs ")

    return text.encode("cp1252", "replace").decode("cp1252")


def to_pdf(data):

    from PIL import Image as PILImage
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    ink, grey = colors.HexColor("#1d1d1f"), colors.HexColor("#86868b")
    card = colors.HexColor("#f1f5ff")

    base = getSampleStyleSheet()["Normal"]

    def style(name, **kw):
        kw.setdefault("textColor", ink)
        return ParagraphStyle(name, parent=base, **kw)

    s_title = style("t", fontName="Helvetica-Bold", fontSize=24, leading=28, spaceAfter=4)
    s_sub = style("s", fontSize=9.5, leading=13, textColor=grey)
    s_h2 = style("h2", fontName="Helvetica-Bold", fontSize=14, leading=18, spaceBefore=16, spaceAfter=6)
    s_h3 = style("h3", fontName="Helvetica-Bold", fontSize=11, leading=14, spaceBefore=8)
    s_body = style("b", fontSize=10, leading=14)
    s_label = style("l", fontSize=8.5, leading=11, textColor=grey)
    s_value = style("v", fontName="Helvetica-Bold", fontSize=19, leading=23)

    def P(text, st):
        return Paragraph(html_lib.escape(pdf_plain(text), quote=False), st)

    story = [P(data["title"], s_title)]

    meta = [data["created"]]
    if data["file"]:
        meta.insert(0, data["file"])
    meta.append(f"{data['rows_in_view']:,} of {data['rows_total']:,} {data['noun']} in view")
    story.append(P("  |  ".join(meta), s_sub))

    for line in data["filters"]:
        story.append(P(line, s_sub))

    # ---------------- KPI cards
    if data["kpis"]:

        story.append(P("Key numbers", s_h2))
        cells = [[P(k["label"], s_label), P(k["value"], s_value), P(k["sub"], s_label)] for k in data["kpis"]]
        rows = [cells[i:i + 4] for i in range(0, len(cells), 4)]
        rows = [r + [""] * (4 - len(r)) for r in rows]

        table = Table(rows, colWidths=[4.25 * cm] * 4)
        table.setStyle(TableStyle(
            [("VALIGN", (0, 0), (-1, -1), "TOP"), ("GRID", (0, 0), (-1, -1), 4, colors.white),
             ("LEFTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 8),
             ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]
            + [("BACKGROUND", (c, r), (c, r), card)
               for r, row in enumerate(rows) for c, cell in enumerate(row) if cell != ""]
        ))
        story.append(table)

    # ---------------- charts
    if data["charts"]:
        story.append(P("Charts", s_h2))

    for item in data["charts"]:

        block = [P(item["title"], s_h3)]

        if item["subtitle"]:
            block.append(P(item["subtitle"], s_sub))

        result = item["result"]
        png = png_of(item) if result["kind"] in ("line", "groups") and item["type"] != "table" else None

        if png:
            width, height = PILImage.open(io.BytesIO(png)).size
            draw_w = 16.5 * cm
            draw_h = draw_w * height / width
            if draw_h > 11 * cm:
                draw_h, draw_w = 11 * cm, 11 * cm * width / height
            block.append(Image(io.BytesIO(png), width=draw_w, height=draw_h))

        elif result["kind"] == "groups":

            header, rows = rows_of(item)
            body = [header] + [
                [name, number_text(value)] + ([f"{share:.1f}"] if share is not None else [])
                for name, value, share in rows
            ]
            table = Table(
                [[pdf_plain(c) for c in row] for row in body], repeatRows=1,
                colWidths=[8 * cm, 4 * cm] + ([3 * cm] if len(header) == 3 else []),
            )
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), card), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9), ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#d8dde8")),
            ]))
            block.append(table)

        else:
            block.append(P(result.get("message", "No data."), s_body))

        story += [KeepTogether(block), Spacer(1, 6)]

    # ---------------- insights
    if data["insights"]:

        story.append(P("Insights (whole file)", s_h2))

        for number, item in enumerate(data["insights"], start=1):
            story.append(Paragraph(
                f"<b>{number}. {html_lib.escape(pdf_plain(item['title']), quote=False)}</b> "
                f"<font color='#86868b'>({item['tag']})</font><br/>"
                f"{html_lib.escape(pdf_plain(item['text']), quote=False)}",
                s_body,
            ))
            story.append(Spacer(1, 5))

    # ---------------- data quality
    q = data["quality"]
    story.append(P("The data", s_h2))
    story.append(P(
        f"{q['rows']:,} rows and {q['columns']} columns. {q['missing']:,} missing values "
        f"and {q['duplicates']:,} duplicate rows. Every number in this report is "
        f"calculated from the data by code.", s_body))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(grey)
        canvas.drawString(2 * cm, 1.2 * cm, "Datalyst  |  " + pdf_plain(data["title"])[:80])
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Page {doc.page}")
        canvas.restoreState()

    buffer = io.BytesIO()
    SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=1.8 * cm, bottomMargin=2 * cm, title=pdf_plain(data["title"]),
    ).build(story, onFirstPage=footer, onLaterPages=footer)

    return buffer.getvalue()


# ============================================================
# Word
# ============================================================

def to_docx(data):

    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()

    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.9)

    grey = RGBColor(0x86, 0x86, 0x8B)

    def shade(cell, fill):
        properties = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), fill)
        properties.append(shading)

    def small(text):
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(clean(text))
        run.font.size = Pt(9)
        run.font.color.rgb = grey
        paragraph.paragraph_format.space_after = Pt(1)
        return paragraph

    doc.add_heading(clean(data["title"]), 0)

    meta = [data["created"]]
    if data["file"]:
        meta.insert(0, data["file"])
    meta.append(f"{data['rows_in_view']:,} of {data['rows_total']:,} {data['noun']} in view")
    small("  |  ".join(meta))

    for line in data["filters"]:
        small(line)

    if data["kpis"]:

        doc.add_heading("Key numbers", 1)
        table = doc.add_table(rows=0, cols=4)

        for start in range(0, len(data["kpis"]), 4):

            row = table.add_row().cells

            for offset in range(4):

                if start + offset >= len(data["kpis"]):
                    continue

                kpi = data["kpis"][start + offset]
                cell = row[offset]
                shade(cell, "F1F5FF")

                label = cell.paragraphs[0]
                run = label.add_run(clean(kpi["label"]))
                run.font.size, run.font.color.rgb = Pt(9), grey

                value = cell.add_paragraph()
                run = value.add_run(clean(kpi["value"]))
                run.bold, run.font.size = True, Pt(18)

                sub = cell.add_paragraph()
                run = sub.add_run(clean(kpi["sub"]))
                run.font.size, run.font.color.rgb = Pt(8), grey

    if data["charts"]:
        doc.add_heading("Charts", 1)

    for item in data["charts"]:

        doc.add_heading(clean(item["title"]), 2)

        if item["subtitle"]:
            small(item["subtitle"])

        result = item["result"]
        png = png_of(item) if result["kind"] in ("line", "groups") and item["type"] != "table" else None

        if png:
            doc.add_picture(io.BytesIO(png), width=Inches(6.2))

        elif result["kind"] == "groups":

            header, rows = rows_of(item)
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"

            for index, name in enumerate(header):
                table.rows[0].cells[index].text = clean(name)

            for name, value, share in rows:
                cells = table.add_row().cells
                cells[0].text = clean(name)
                cells[1].text = number_text(value)
                if share is not None:
                    cells[2].text = f"{share:.1f}"

        else:
            doc.add_paragraph(clean(result.get("message", "No data.")))

    if data["insights"]:

        doc.add_heading("Insights (whole file)", 1)

        for number, item in enumerate(data["insights"], start=1):
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(f"{number}. {clean(item['title'])}")
            run.bold = True
            paragraph.add_run(f"  ({item['tag']})").font.color.rgb = grey
            doc.add_paragraph(clean(item["text"]))

    q = data["quality"]
    doc.add_heading("The data", 1)
    doc.add_paragraph(
        f"{q['rows']:,} rows and {q['columns']} columns. {q['missing']:,} missing "
        f"values and {q['duplicates']:,} duplicate rows. Every number in this "
        f"report is calculated from the data by code."
    )

    buffer = io.BytesIO()
    doc.save(buffer)

    return buffer.getvalue()


# ============================================================
# Excel
# ============================================================

def to_xlsx(data):

    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill

    book = Workbook()
    head_fill = PatternFill("solid", fgColor="E6EEFF")
    bold = Font(bold=True)
    used = set()

    def text(sheet, row, column, value, font=None):
        """Always stored as TEXT, so a value like '=1+1' can never run as a formula."""
        cell = sheet.cell(row=row, column=column)
        cell.value = clean(value)
        cell.data_type = "s"
        if font:
            cell.font = font
        return cell

    def sheet_name(name):
        base = re.sub(r"[\[\]:*?/\\]", " ", clean(name)).strip()[:28] or "Sheet"
        candidate, number = base, 1
        while candidate.lower() in used:
            number += 1
            candidate = f"{base[:25]} {number}"
        used.add(candidate.lower())
        return candidate

    summary = book.active
    summary.title = sheet_name("Summary")
    text(summary, 1, 1, data["title"], Font(bold=True, size=16))
    row = 3
    for line in [f"File: {data['file']}" if data["file"] else None, f"Created: {data['created']}",
                 f"Rows in view: {data['rows_in_view']:,} of {data['rows_total']:,} {data['noun']}"] + data["filters"]:
        if line:
            text(summary, row, 1, line)
            row += 1
    q = data["quality"]
    row += 1
    text(summary, row, 1, "The data", bold)
    for line in [f"{q['rows']:,} rows, {q['columns']} columns", f"{q['missing']:,} missing values",
                 f"{q['duplicates']:,} duplicate rows"]:
        row += 1
        text(summary, row, 1, line)
    summary.column_dimensions["A"].width = 70

    kpis = book.create_sheet(sheet_name("KPIs"))
    for column, name in enumerate(["KPI", "Value", "Note"], start=1):
        text(kpis, 1, column, name, bold).fill = head_fill
    for number, kpi in enumerate(data["kpis"], start=2):
        text(kpis, number, 1, kpi["label"])
        value = kpis.cell(row=number, column=2)
        value.value = kpi["raw"]
        value.number_format = "#,##0.00"
        text(kpis, number, 3, kpi["sub"])
    kpis.column_dimensions["A"].width = 38
    kpis.column_dimensions["B"].width = 18
    kpis.column_dimensions["C"].width = 40

    for number, item in enumerate(data["charts"], start=1):

        header, rows = rows_of(item)
        sheet = book.create_sheet(sheet_name(f"{number} {item['title']}"))
        text(sheet, 1, 1, item["title"], Font(bold=True, size=13))

        if not header:
            text(sheet, 3, 1, item["result"].get("message", "No data."))
            continue

        for column, name in enumerate(header, start=1):
            text(sheet, 3, column, name, bold).fill = head_fill

        for offset, (name, value, share) in enumerate(rows, start=4):
            text(sheet, offset, 1, name)
            sheet.cell(row=offset, column=2).value = float(value)
            sheet.cell(row=offset, column=2).number_format = "#,##0.00"
            if share is not None:
                sheet.cell(row=offset, column=3).value = float(share)

        sheet.column_dimensions["A"].width = 32
        sheet.column_dimensions["B"].width = 18

        last = 3 + len(rows)

        if rows:

            if item["type"] == "line":
                chart = LineChart()
            elif item["type"] == "donut":
                chart = PieChart()
            else:
                chart = BarChart()
                chart.type = "col" if item["type"] == "column" else "bar"

            chart.title = clean(item["title"])[:80]
            chart.add_data(Reference(sheet, min_col=2, min_row=3, max_row=last), titles_from_data=True)
            chart.set_categories(Reference(sheet, min_col=1, min_row=4, max_row=last))
            chart.height, chart.width = 8, 16
            sheet.add_chart(chart, "E3")

    if data["insights"]:
        sheet = book.create_sheet(sheet_name("Insights"))
        for column, name in enumerate(["Type", "Title", "Finding"], start=1):
            text(sheet, 1, column, name, bold).fill = head_fill
        for number, item in enumerate(data["insights"], start=2):
            text(sheet, number, 1, item["tag"])
            text(sheet, number, 2, item["title"])
            cell = text(sheet, number, 3, item["text"])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.column_dimensions["B"].width = 30
        sheet.column_dimensions["C"].width = 100

    buffer = io.BytesIO()
    book.save(buffer)

    return buffer.getvalue()


# ============================================================
# Web page (one file, works offline, prints well)
# ============================================================

PAGE_STYLE = """
body{font-family:-apple-system,BlinkMacSystemFont,'SF Pro Text','Segoe UI',Helvetica,Arial,sans-serif;background:#f4f6fb;color:#1d1d1f;margin:0;padding:32px 20px;}
.wrap{max-width:1000px;margin:0 auto}
h1{font-size:34px;letter-spacing:-.03em;margin:0 0 6px}
.meta{color:#86868b;font-size:13px;line-height:1.6}
h2{font-size:20px;margin:30px 0 12px;letter-spacing:-.02em}
.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(200px,1fr));gap:14px}
.kpi,.card{background:#fff;border-radius:18px;padding:18px 20px;box-shadow:0 6px 24px rgba(60,72,120,.09);}
.kpi .l{font-size:12.5px;color:#86868b;font-weight:600}.kpi .v{font-size:28px;font-weight:750;letter-spacing:-.03em;margin:4px 0}.kpi .s{font-size:12px;color:#86868b}
.card{margin-bottom:16px;page-break-inside:avoid}.card h3{margin:0 0 2px;font-size:16px}.card .sub{color:#86868b;font-size:12.5px;margin:0 0 8px}
.card svg{max-width:100%;height:auto}
table{border-collapse:collapse;width:100%;font-size:13.5px}th{background:#eef3ff;text-align:left}th,td{padding:7px 10px;border-bottom:1px solid #e3e8f3}td.n,th.n{text-align:right}
.ins{background:#fff;border-radius:14px;padding:14px 18px;margin-bottom:10px;box-shadow:0 4px 16px rgba(60,72,120,.07)}
.tag{font-size:11px;font-weight:700;letter-spacing:.07em;text-transform:uppercase;color:#0a6cf5}
.foot{color:#86868b;font-size:12px;margin-top:30px}
@media print{body{background:#fff;padding:0}.kpi,.card,.ins{box-shadow:none;border:1px solid #e3e8f3}}
"""


def to_html(data):

    esc = lambda value: html_lib.escape(clean(value))

    parts = ["<!DOCTYPE html><html lang='en'><head><meta charset='utf-8'>",
             "<meta name='viewport' content='width=device-width, initial-scale=1'>",
             f"<title>{esc(data['title'])}</title><style>{PAGE_STYLE}</style></head><body><div class='wrap'>",
             f"<h1>{esc(data['title'])}</h1><div class='meta'>"]

    meta = [data["created"]]
    if data["file"]:
        meta.insert(0, data["file"])
    meta.append(f"{data['rows_in_view']:,} of {data['rows_total']:,} {data['noun']} in view")
    parts.append(" &middot; ".join(esc(m) for m in meta))
    parts += [f"<br>{esc(line)}" for line in data["filters"]]
    parts.append("</div>")

    if data["kpis"]:
        parts.append("<h2>Key numbers</h2><div class='kpis'>")
        parts += [
            f"<div class='kpi'><div class='l'>{esc(k['label'])}</div><div class='v'>{esc(k['value'])}</div>"
            f"<div class='s'>{esc(k['sub'])}</div></div>" for k in data["kpis"]
        ]
        parts.append("</div>")

    if data["charts"]:
        parts.append("<h2>Charts</h2>")

    for number, item in enumerate(data["charts"], start=1):

        parts.append(f"<div class='card'><h3>{esc(item['title'])}</h3><p class='sub'>{esc(item['subtitle'])}</p>")

        result = item["result"]
        svg = svg_of(item, number) if result["kind"] in ("line", "groups") and item["type"] != "table" else None

        if svg:
            parts.append(svg)

        elif result["kind"] == "groups":

            header, rows = rows_of(item)
            parts.append("<table><tr>" + "".join(
                f"<th class='{'n' if i else ''}'>{esc(h)}</th>" for i, h in enumerate(header)) + "</tr>")
            for name, value, share in rows:
                parts.append(f"<tr><td>{esc(name)}</td><td class='n'>{esc(number_text(value))}</td>"
                             + (f"<td class='n'>{share:.1f}</td>" if share is not None else "") + "</tr>")
            parts.append("</table>")

        else:
            parts.append(f"<p>{esc(result.get('message', 'No data.'))}</p>")

        parts.append("</div>")

    if data["insights"]:
        parts.append("<h2>Insights (whole file)</h2>")
        parts += [
            f"<div class='ins'><span class='tag'>{esc(i['tag'])}</span><br><b>{esc(i['title'])}</b><br>{esc(i['text'])}</div>"
            for i in data["insights"]
        ]

    q = data["quality"]
    parts.append(
        f"<div class='foot'>{q['rows']:,} rows and {q['columns']} columns. {q['missing']:,} missing values and "
        f"{q['duplicates']:,} duplicate rows. Every number in this report is calculated from the data by code. "
        f"Made with Datalyst.</div></div></body></html>"
    )

    return "".join(parts)


EXPORTERS = {"pdf": to_pdf, "docx": to_docx, "xlsx": to_xlsx, "html": to_html}


def export(data, formats):
    """{format: bytes} for the chosen formats."""

    files = {}

    for key in formats:
        output = EXPORTERS[key](data)
        files[key] = output.encode("utf-8") if isinstance(output, str) else output

    return files
