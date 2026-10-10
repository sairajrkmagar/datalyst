"""
REPORT EXPORT TEST (no AI needed)

For every dataset it builds a report (PDF, Word, Excel, web page) and checks that:
  - the KPI numbers equal an independent pandas calculation,
  - every file carries exactly the numbers of the dashboard,
  - every file can be opened again,
  - hostile text (<script>, =1+1, a rupee sign, an emoji, control characters)
    can neither break a file nor run as code,
  - empty dashboards, empty filters and broken charts still give a report.

    python test_report.py            (all datasets)
    python test_report.py amazon.csv
"""
import io
import random
import re
import sys
import warnings

warnings.filterwarnings("ignore")

import pandas as pd

import builder as b
import report as r
from ask_engine import AskEngine

missing = r.missing_libraries()
if missing:
    print("SKIPPED: missing libraries. Install them with:  pip install " + " ".join(missing))
    sys.exit(0)

import docx
import openpyxl

try:
    import pymupdf as pdf_reader
except Exception:
    try:
        import fitz as pdf_reader
    except Exception:
        pdf_reader = None

import test_builder as T

FILES = ["ecommerce_sales_34500.csv", "sample_sales.csv", "Cleaned_HR_Data_Analysis.csv", "amazon.csv"]
TRICKY = "=1+1 <script>alert(1)</script> & \"quotes\" ₹ 😀 \x07 end"

passed, failed = 0, []


def check(ok, message):
    global passed
    if ok:
        passed += 1
    else:
        failed.append(message)


def expected_kpi(view, tile):
    if len(view) == 0:
        return None
    m, agg = tile["measure"], tile["agg"]
    if agg == "count":
        return float(len(view))
    if agg == "distinct":
        return float(view[m].dropna().nunique())
    s = pd.to_numeric(view[m], errors="coerce").dropna()
    return None if s.empty else float(getattr(s, agg)())


def close(a, c):
    return abs(a - c) <= 1e-6 * max(1.0, abs(c))


def images_in(data):
    return sum(1 for c in data["charts"] if c["result"]["kind"] in ("line", "groups") and c["type"] != "table")


def docx_text(document):
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts += [p.text for p in cell.paragraphs]
    return "\n".join(parts)


def verify(label, info, view, tiles, data, files, opts):
    kpi_tiles = [t for t in tiles if t["type"] == "kpi"]
    clean_title = r.clean(data["title"])

    # ---- the numbers vs an independent calculation
    check(len(data["kpis"]) == len(kpi_tiles), f"{label}: KPI count")
    for tile, kpi in zip(kpi_tiles, data["kpis"]):
        want = expected_kpi(view, tile)
        same = (want is None and kpi["raw"] is None) or (want is not None and kpi["raw"] is not None and close(kpi["raw"], want))
        check(same, f"{label}: KPI {kpi['label']} {kpi['raw']} vs {want}")

    # ---- Excel: opens, numbers equal, nothing runs as a formula
    book = openpyxl.load_workbook(io.BytesIO(files["xlsx"]))
    sheet = book["KPIs"]
    for number, kpi in enumerate(data["kpis"], start=2):
        value = sheet.cell(row=number, column=2).value
        check((value is None and kpi["raw"] is None) or (value is not None and kpi["raw"] is not None and close(value, kpi["raw"])), f"{label}: Excel KPI row {number}")
    chart_sheets = [s for s in book.sheetnames if s[:1].isdigit()]
    check(len(chart_sheets) == len(data["charts"]), f"{label}: Excel has {len(chart_sheets)} chart sheets, expected {len(data['charts'])}")
    chart_tiles = [t_ for t_ in tiles if t_["type"] != "kpi"]
    for name, tile in zip(chart_sheets, chart_tiles):
        # a FRESH calculation, not the exporter's own helper
        fresh = b.compute_tile(view, tile, info, opts)
        expected = list(fresh["values"].items()) if fresh["kind"] == "groups" else (list(fresh["months"].items()) if fresh["kind"] == "line" else [])
        ws = book[name]
        for offset, (key, value) in enumerate(expected, start=4):
            check(str(ws.cell(row=offset, column=1).value) == r.clean(key), f"{label}: Excel {name} name row {offset}")
            check(close(ws.cell(row=offset, column=2).value, value), f"{label}: Excel {name} value row {offset}")
        check(ws.cell(row=4 + len(expected), column=1).value is None, f"{label}: Excel {name} has extra rows")
    summary = book[book.sheetnames[0]]
    check(summary["A1"].data_type == "s" and "=1+1" in str(summary["A1"].value), f"{label}: Excel title must stay text")
    check(all(c.data_type != "f" for ws in book.worksheets for row in ws.iter_rows() for c in row), f"{label}: Excel contains a formula")

    # ---- Word: opens, text and pictures
    document = docx.Document(io.BytesIO(files["docx"]))
    text = docx_text(document)
    check(clean_title in text, f"{label}: Word title")
    for kpi in data["kpis"]:
        check(kpi["label"] in text and kpi["value"] in text, f"{label}: Word KPI {kpi['label']}")
    for item in data["charts"]:
        check(r.clean(item["title"]) in text, f"{label}: Word chart title {item['title']}")
    for item in data["insights"]:
        check(r.clean(item["title"]) in text, f"{label}: Word insight {item['title']}")
    check(len(document.inline_shapes) == images_in(data), f"{label}: Word pictures {len(document.inline_shapes)} vs {images_in(data)}")

    # ---- PDF: opens, text present
    check(files["pdf"][:5] == b"%PDF-", f"{label}: PDF header")
    if pdf_reader:
        pdf = pdf_reader.open(stream=files["pdf"], filetype="pdf")
        check(len(pdf) >= 1, f"{label}: PDF pages")
        pdf_text = "\n".join(page.get_text() for page in pdf).replace("\n", " ")
        pdf_text = re.sub(r"\s+", " ", pdf_text)
        for kpi in data["kpis"]:
            check(re.sub(r"\s+", " ", r.pdf_plain(kpi["label"])) in pdf_text, f"{label}: PDF KPI label {kpi['label']}")
            check(kpi["value"] in pdf_text, f"{label}: PDF KPI value {kpi['value']}")
        for item in data["charts"]:
            check(re.sub(r"\s+", " ", r.pdf_plain(item["title"]))[:30] in pdf_text, f"{label}: PDF chart title {item['title']}")
        check("Rs" in pdf_text, f"{label}: PDF rupee sign replaced")

    # ---- Web page: one file, no outside links, nothing executable
    page = files["html"].decode("utf-8")
    check("<script" not in page.lower(), f"{label}: web page contains a script tag")
    check(not re.search(r"""(src|href)=["']https?:""", page, re.I), f"{label}: web page links to the outside")
    check(page.count("<svg") == images_in(data), f"{label}: web page charts {page.count('<svg')} vs {images_in(data)}")
    check("&lt;script&gt;" in page, f"{label}: web page did not escape the title")
    for kpi in data["kpis"]:
        check(kpi["label"].replace("&", "&amp;") in page and kpi["value"] in page, f"{label}: web page KPI {kpi['label']}")
    ids = re.findall(r'\bid="([^"]+)"', page)
    check(len(ids) == len(set(ids)), f"{label}: web page has duplicate svg ids")


def run(path, seed):
    df = pd.read_csv(path)
    engine = AskEngine(df)
    info = b.dataset_info(engine)
    options = b.filter_options(info)
    rng = random.Random(seed)

    tiles = b.starter_tiles(info)
    if info["chart_dims"]:
        tiles.append(b.new_tile(info, type="table", measure=info["main_measure"], agg="sum" if info["main_measure"] else "count", group=info["chart_dims"][0]))
        tiles.append(b.new_tile(info, type="column", measure=None, agg="count", group=info["chart_dims"][0]))
        tiles.append(b.new_tile(info, type="donut", measure=None, agg="count", group=info["chart_dims"][0], title=TRICKY))

    scenarios = [({}, {"grain": "M", "top": None}),
                 (T.random_filters(options, rng), {"grain": "Q", "top": 5}),
                 (T.random_filters(options, rng), {"grain": "Y", "top": 3})]

    for number, (filters, opts) in enumerate(scenarios):
        view = b.apply_filters(info, filters)
        data = r.build_report(engine, info, tiles, filters, opts, TRICKY, path)
        files = r.export(data, ["pdf", "docx", "xlsx", "html"])
        verify(f"{path} scenario {number}", info, view, tiles, data, files, opts)

    # ---- nothing to show: an empty dashboard, everything filtered out, a broken chart
    nothing = {"values": {d: [] for d in list(options["dims"])[:1]}}
    broken = dict(b.new_tile(info, type="bar", measure=None, agg="count", group=info["dims"][0] if info["dims"] else None), group="no_such_column")
    for label, the_tiles, the_filters in [("empty dashboard", [], {}), ("everything filtered out", tiles, nothing), ("broken chart", tiles + [broken], {})]:
        try:
            data = r.build_report(engine, info, the_tiles, the_filters, {"grain": "M", "top": None}, "x", path)
            files = r.export(data, ["pdf", "docx", "xlsx", "html"])
            check(all(len(v) > 500 for v in files.values()), f"{path}: {label} gave an empty file")
        except Exception as error:
            check(False, f"{path}: {label} crashed: {error!r}")

    print(f"  {path:34s} done")


if __name__ == "__main__":
    # file names are safe
    for title in [TRICKY, "../../etc/passwd", "a/b\\c:d*e?", "", "   ", "x" * 300]:
        name = r.file_name(title, "pdf")
        check(name.endswith(".pdf") and not re.search(r"[/\\:*?<>|\x00]", name) and len(name) < 70, f"file name for {title[:20]!r}: {name!r}")

    files = sys.argv[1:] or FILES
    print("Report export test")
    for i, f in enumerate(files):
        run(f, 300 + i)

    print(f"\n{passed + len(failed)} checks, {len(failed)} FAILED" + ("" if pdf_reader else "   (PDF text not checked: install pymupdf)"))
    for f in failed[:12]:
        print("  FAIL:", f[:240])
